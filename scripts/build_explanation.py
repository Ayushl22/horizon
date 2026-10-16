"""Build the learning guide after UI verification, using saved results and real screenshots.

Run with python-docx/Pillow from the Codex bundled document runtime. This is an
authoring utility, not a dependency of the Streamlit app.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT/'reports'
R = json.loads((ROOT/'artifacts/results.json').read_text(encoding='utf-8'))
A = R['audit']
M = R['models']
SHOTS = REPORTS/'ui_screenshots'
OUT = REPORTS/'Customer_Churn_Project_Explanation.docx'
TITLES = [
 'Problem and end to end workflow', 'Technology and file organisation',
 'Dataset and input feature dictionary', 'Data cleaning and the seven charge records',
 'Splitting stratification and leakage prevention', 'Preprocessing and consistent inference',
 'Logistic Regression and Random Forest', 'Decision Tree and K Means segmentation',
 'Training model selection and saved artifacts', 'Evaluation metrics with a worked example',
 'Actual saved experimental results', 'Using the Overview page', 'Using the single customer form',
 'Reading predictions and exploring what if inputs', 'Using batch prediction and CSV export',
 'Using the Model Comparison page', 'Using Customer Insights',
 'Feature importance and retention limitations', 'Using the customer segmentation view', 'Tests and completed verification',
 'Setup running and deployment', 'Limitations and future improvements',
 'Viva questions and answers', 'Sources and project references',
]

doc = Document()
sec = doc.sections[0]
sec.page_width,sec.page_height = Inches(8.5),Inches(11)
sec.top_margin,sec.bottom_margin = Inches(.65),Inches(.65)
sec.left_margin,sec.right_margin = Inches(.75),Inches(.75)
sec.footer_distance = Inches(.25)
for name in ['Normal','Body Text']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(11); s.font.bold=False
    s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.1
for name,size in [('Title',28),('Subtitle',14),('Heading 1',20),('Heading 2',13),('Heading 3',11)]:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size)
    s.font.bold=name!='Subtitle'; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(12 if name!='Title' else 0)
    s.paragraph_format.space_after=Pt(8)
    s.paragraph_format.keep_with_next=True
for name in ['Caption','List Bullet','List Number']:
    doc.styles[name].font.name='Calibri'; doc.styles[name].font.color.rgb=RGBColor(0,0,0)
# The bundled Word starter has a blue border on Title; remove inherited residue.
for node in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
doc.styles['Caption'].font.size=Pt(10)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); run._r.addnext(field)
doc.core_properties.title='Customer Churn Project Explanation'
doc.core_properties.subject='Learning and viva guide to the implemented customer retention dashboard'
doc.core_properties.author='Project documentation'

def p(text='',style=None):
    return doc.add_paragraph(text,style)

def h(text):
    doc.add_heading(text,2)

def bullet(text):
    p(text,'List Bullet')

def chapter(index):
    paragraph=doc.add_heading(TITLES[index],1)
    # Break on the heading itself: a standalone break paragraph can spill onto
    # a blank page when the previous chapter ends near the bottom margin.
    paragraph.paragraph_format.page_break_before=True
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(index+1)); start.set(qn('w:name'),f'chapter_{index}')
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(index+1))
    paragraph._p.insert(0,start); paragraph._p.append(end)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for cell,width in zip(t.columns,widths):
        cell.width=Inches(width)
    for row in [t.rows[0]]:
        repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
    for cell,text in zip(t.rows[0].cells,headers): cell.text=str(text)
    for values in rows:
        for cell,text in zip(t.add_row().cells,values): cell.text=str(text)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
            props.append(borders)
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'17354D' if i==0 else ('F1F5F8' if i%2==0 else 'FFFFFF')); props.append(shading)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                val=OxmlElement('w:'+side); val.set(qn('w:w'),'90'); val.set(qn('w:type'),'dxa'); margins.append(val)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2); paragraph.paragraph_format.line_spacing=1.05
                if j>0 and all(len(str(r[j]))<15 for r in rows): paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size=Pt(10.5); run.font.bold=i==0
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
            no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    p('')
    return t

def picture(filename,caption,width=7):
    paragraph=doc.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run(); shape=run.add_picture(str(SHOTS/filename),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    paragraph.paragraph_format.space_after=Pt(5)
    p(caption,'Caption')

def code(lines):
    for line in lines:
        para=p(line); para.paragraph_format.space_after=Pt(3)
        for run in para.runs: run.font.name='Consolas'; run.font.size=Pt(9.5)

# Cover establishes the audience and key conclusion without a decorative title rule.
p('Customer Churn Prediction and Retention Dashboard','Title')
p('Project explanation and viva guide','Subtitle')
p('A practical learning guide for a second year CSE student')
h('What you will learn')
p('This guide explains the implemented Python project from the source CSV to the final prediction screen. It connects the actual code, cleaning decisions, trained models, evaluation results and redesigned Streamlit pages so you can demonstrate the application and explain it in a viva.')
p('The revised project has three supervised churn classifiers and one unsupervised segmentation model. Random Forest wins classifier selection using validation F1. K-Means is chosen separately using training feature geometry. Its cluster numbers are arbitrary; silhouette cannot be ranked against classifier F1. Instructor approval may be needed if the rubric requires four supervised or state-of-the-art classifiers; full compliance with that interpretation is not claimed.')
table(['Project fact','Saved implementation'],[
    ['Dataset','UCI Iranian Churn, 3,150 original records'],['Analysed profiles',f'{A["clean_rows"]:,} after documented cleaning'],
    ['Models','Logistic Regression, Random Forest, Decision Tree and K-Means'],['Selected model',R['selected_model']],
    ['Decision threshold',f'{R["threshold"]:.2f}'],['Interface','Five Streamlit pages using saved model pipelines'],
    ['Guide date','16 October 2026']], [2.1,4.9])
p('Read the workflow and data chapters first, then the models and metrics. Use the UI chapters beside the running application. The final questions provide a short revision exercise. This guide is separate from the preserved ten-page submission PDF.')

doc.add_page_break()
doc.add_heading('Contents',1)
p('Each entry links to its chapter. Page numbers refer to this generated edition.')
for i,title in enumerate(TITLES):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(5)
    hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('w:anchor'),f'chapter_{i}')
    r=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text=f'{i+1:02d}  {title}'; r.append(text); hyperlink.append(r); para._p.append(hyperlink)
    para.paragraph_format.tab_stops.add_tab_stop(Inches(6.7),WD_TAB_ALIGNMENT.RIGHT)
    mapping_path=ROOT/'tmp/docx_review/page_map.json'
    page_map=json.loads(mapping_path.read_text()) if mapping_path.exists() else {}
    para.add_run(f'\t{page_map.get(title,i+3)}')
p('Screenshots show the actual final local application. Some views are scrolled to the relevant result or chart. Local browser checks do not imply that cloud deployment or physical-phone testing has been completed.')

chapter(0)
h('The problem in simple terms')
p('Customer churn means a customer stops using a company’s service. For example, a telecom subscriber may move to another operator or stop renewing. A provider would like an early warning so staff can review the situation before the customer leaves. A churn model estimates which profiles resemble customers who later churned in its training data.')
p('In this dataset, the predictors summarise the first nine months. The label records churn at the end of month twelve. UCI describes the intervening three months as a planning gap. The project cannot reconstruct that timing independently because the CSV has no original event timestamps.')
h('The implemented objective')
p('Compare three modest CPU classifiers and a separate K-Means segmentation model using one reproducible split, select a model using validation F1, report test results honestly, and let a user score either one customer or an unlabelled CSV. The dashboard is a decision-support demonstration; it does not automatically contact customers or promise that any offer will retain them.')
h('Follow one record through the project')
for text in [
 'Download the official ZIP into data/raw and load the real CSV. Normalize header whitespace and check the documented schema.',
 'Quarantine seven charge-band anomalies, remove conflicting retained-feature profiles, and collapse repeated profiles.',
 'Make stratified train, validation and test partitions. Fit preprocessing and all estimator parameters on training only.',
 'Score validation data, select the highest churn F1 and freeze that choice. Then calculate final test metrics and save pipelines and results.',
 'The UI validates new inputs, selects the stored feature order, calls the saved pipeline and displays probability plus the thresholded class.']:
    p(text,'List Number')
h('Two separate paths')
p('Training is a command-line task in src/train.py. Inference is a dashboard task using src/predict.py. Opening a page or changing a widget does not train a new model. That keeps the demonstration fast and ensures that the UI uses the same evaluated model.')

chapter(1)
table(['Library','Purpose in this project'],[
 ['Python','Application language and command-line modules'],['pandas and NumPy','Tables, numeric conversion, validation masks and arrays'],
 ['scikit-learn','Splits, preprocessing, Logistic Regression, Random Forest, metrics and permutation importance'],
 ['scikit-learn K-Means and PCA','Feature-distance segmentation and a training-fitted 2D projection'],['Streamlit','Navigation, forms, upload, tables, status messages and downloads'],
 ['Plotly / Matplotlib','Interactive dashboard charts / saved report chart PNGs'],['joblib','Persist and load fitted pipelines with metadata'],
 ['ReportLab and pypdf','Generate and count pages in the existing ten-page PDF'],['pytest','Automated regression tests']], [1.7,5.3])
h('Where to read the code')
table(['Location','Responsibility'],[
 ['app.py','Configure Streamlit, load saved results and route the chosen page'],['src/config.py','Root paths, seed, exact feature order, schema and exclusions'],
 ['src/data.py / preprocessing.py','Official data loading, cleaning, splits and transform definition'],['src/train.py / evaluate.py','Fit three classifiers, tune Decision Tree on validation and report metrics'],
 ['src/predict.py / segmentation.py','Churn inference, K-Means fit and saved cluster assignment'],['src/ui.py / ui_pages.py','Shared UI components and all five page layouts'],
 ['src/ui_charts.py / ui.css','Consistent chart colours and small responsive stylesheet'],['artifacts / data/raw','Fitted models and saved outputs / original attributed CSV'],
 ['tests / reports / .streamlit','Checks / documentation and screenshots / app configuration']], [2.4,4.6])
p('The project has no database, GPU service or paid API. The Word authoring utility is separate from app dependencies. Paths are anchored to the project root in config.py so modules write outputs consistently.')

chapter(2)
p('Source: Iranian Churn [Dataset] (2020), UCI Machine Learning Repository, DOI 10.24432/C5JW3Z, licensed CC BY 4.0. Attribution and the official source link are retained in data/raw/ATTRIBUTION.md and the dashboard sidebar. The source reports 3,150 customers and 13 predictors; the CSV has 14 columns including the target.')
table(['Exact CSV feature','Meaning and accepted input'],[
 ['Call Failure','Number of failed calls; integer at least 0'],['Complains','0 = no complaint; 1 = complaint'],
 ['Subscription Length','Total subscription duration in months; integer at least 0'],['Charge Amount','Ordinal charge band 0 lowest to 9 highest; not a currency amount'],
 ['Seconds of Use','Total call duration in seconds; integer at least 0'],['Frequency of use','Total number of calls; integer at least 0'],
 ['Frequency of SMS','Total number of text messages; integer at least 0'],['Distinct Called Numbers','Number of distinct called numbers; integer at least 0'],
 ['Age Group','Ordinal band 1 youngest to 5 oldest; exact age cutoffs are not supplied'],['Tariff Plan','1 = pay as you go; 2 = contractual']], [2.2,4.8])
h('Target and extra columns')
p('Churn is the target: 1 means churn and 0 means non-churn. Prediction files do not need this column because the model estimates it. The source also contains Status, Age and Customer Value; the next chapter explains their exclusion. No reliable customer identifier appears in the official CSV.')
p('Human-readable form labels map back to these exact headers in src/ui.py. For example, Failed calls becomes Call Failure. A batch CSV must retain the sample headers. An extra customer_id column is preserved in the exported results but is never used as a predictive feature.')
p('Counts have no documented maximum. The code rejects negative or fractional counts rather than imposing the largest training value as a rule. It does not delete a customer merely because their valid usage is unusually high.')

chapter(3)
table(['Audit item','Actual saved finding'],[
 ['Raw records',A['raw_rows']],['Missing cells',sum(A['missing'].values())],['Exact raw duplicate rows',A['exact_duplicates']],
 ['Charge Amount 10 records',len(A['quarantined_charge_rows'])],['Eligible conflicting-profile rows',A['conflicting_profile_rows']],
 ['All rows removed',A['removed_rows']],['Retained distinct profiles',A['clean_rows']]], [4.2,2.8])
h('Why seven charge records are excluded')
p('The official variable notes define Charge Amount as an ordinal band from 0 to 9. Seven rows in the downloaded official CSV instead contain 10. Their zero-based source indices are '+', '.join(map(str,A['quarantined_charge_rows']))+'. This is a real conflict between the supplied file and its documentation, rather than a value we introduced.')
p('The implementation preserves the raw CSV and quarantines those rows before the experiment. It does not silently relabel 10 as 9 or enlarge the prediction schema. The evidence does not establish whether 10 is a data-entry error or an undocumented valid band. Exclusion is a conservative choice; checking with the dataset provider or doing a separate sensitivity analysis would resolve that uncertainty more convincingly.')
h('Profiles and conflicting labels')
p('After quarantine, records are grouped by the ten retained features. All rows of a profile with both churn labels are removed: a deterministic classifier cannot distinguish those identical inputs. Remaining repeats are collapsed. This prevents identical retained profiles appearing in different partitions, but may merge distinct people and changes the population. Raw duplicate counts overlap these operations; do not add every audit count as though the groups were disjoint.')
h('Features excluded before modelling')
p('Status codes active/non-active accounts. It is excluded conservatively as an already-inactive proxy despite UCI’s timing statement. Customer Value has an unpublished formula, so its ingredients and timing cannot be audited. Age is not separately explained in the UCI notes, so the documented Age Group is used instead. Identifiers are excluded. No missing values or invalid outliers are invented to demonstrate cleaning.')

chapter(4)
table(['Partition','Profiles','Churn','Purpose'],[[name.title(),v['rows'],v['churn'],{'train':'Fit preprocessing and model parameters','validation':'Compare models and select the winner','test':'Final reporting after selection'}[name]] for name,v in A['splits'].items()],[1.2,.9,.9,4])
p('src/data.py first takes 30% as a holdout, stratifying by Churn with seed 42. It then divides that holdout equally into validation and test, again stratifying. Rounded integer counts give about 70/15/15. Stratification keeps the minority class represented in each partition rather than relying on an accidental random composition.')
h('Why the test set is separate')
p('Imagine repeatedly checking an exam answer sheet while choosing your solution. You would gradually optimise for that sheet rather than demonstrate what you know. Similarly, choosing a model or threshold after examining its test results makes the test less independent. Here validation is the selection workspace; test is final reporting only.')
h('Leakage protections in the actual code')
bullet('Cleaning isolates exact retained-feature profiles before splitting. Saved source indices make the assignment auditable; the same split serves all four models.')
bullet('Imputers, scaler, encoder and model parameters are fitted on training data only. Validation/test values do not determine training transformations.')
bullet('Weights for class imbalance come from training. No resampling occurs, either before or after splitting.')
bullet('A fixed 0.5 threshold and validation F1 selection are used. Test scores never alter the winner, threshold or parameters.')
bullet('Status, Customer Value and Age are excluded for the reasons already documented.')
p('A remaining limitation is the absence of customer IDs and original timestamps. Exact profile isolation cannot prove that two different rows never represent the same person, and the claimed observation timeline cannot be independently verified. A random split is also not a substitute for testing on a later time period.')

chapter(5)
h('The preprocessing pipeline')
p('make_preprocessor in src/preprocessing.py returns a ColumnTransformer with two branches. Numeric features are all retained columns except Complains and Tariff Plan. A median SimpleImputer handles missing training values and StandardScaler subtracts the training mean and divides by the training scale. The source has no missing cells, so imputation is a defensive capability rather than a fabricated cleaning step.')
p('Complains and Tariff Plan use a most-frequent imputer and a OneHotEncoder. The encoder produces a dense numeric matrix and ignores unknown categories internally. Public prediction validation is stricter and rejects undocumented categories before they reach it. Charge Amount and Age Group retain their ordinal numeric meaning.')
h('Why scaling matters')
p('Call duration can be thousands of seconds while a charge band is a single digit. Scaling prevents raw units from dominating a regularised linear fit. Trees generally do not need scaling, but this project deliberately uses the same preprocessing definition for a simple, consistent comparison.')
h('How inference stays consistent')
for text in ['load_model reads a trusted joblib bundle and checks feature order and key package versions.',
 'validate_features checks required columns, numeric conversion, finite numbers, category membership and nonnegative integer counts.',
 'predict selects only the ten known features in their saved order. Extra metadata, IDs and any supplied target are excluded.',
 'pipeline.predict_proba applies the already fitted transformations and returns the probability for class 1. The code never calls fit on an uploaded file.',
 'Probability at least 0.5 becomes predicted_churn = 1. The output retains the input columns and adds probability and class.']:
    p(text,'List Number')
p('The UI rejects missing prediction values so the user can correct the file rather than silently rely on imputation. Save the entire pipeline, not just the estimator: otherwise a form may feed raw values into a model that expects transformed columns.')

chapter(6)
h('Logistic Regression as the baseline')
p('Logistic Regression combines transformed inputs into a weighted score and passes it through a sigmoid function to estimate a binary-class probability. Despite its name, this is a classification model. Its decision surface is linear in the transformed features, which makes it a useful inexpensive baseline.')
table(['Important setting','Actual value','Why it matters'],[
 ['max_iter','1500','Allows the numerical solver enough iterations'],['class_weight','balanced','Gives the minority churn class more influence'],['random_state','42','Records the reproducibility seed']], [1.6,1.2,4.2])
p('Its strengths are fast CPU fitting, a simple parameter structure and a meaningful comparison point. Its limits include weak representation of nonlinear interactions unless we engineer them. A larger coefficient is not proof of a causal effect, and its weighted probability should not be assumed calibrated. The project reports permutation importance consistently across models instead of presenting raw coefficients as causes.')
h('Random Forest as the bagging model')
p('A decision tree asks a sequence of questions such as whether usage is below a split value. A forest fits many trees with randomised data/feature choices and averages their predictions. That averaging can reduce the instability of a single tree and represent interactions without manually writing every rule.')
table(['Important setting','Actual value','Why it matters'],[
 ['n_estimators','180','Number of trees'],['max_depth','12','Bounds tree complexity'],['min_samples_leaf','2','Avoids leaves based on only one training example'],
 ['class_weight','balanced','Weights the churn minority'],['n_jobs / seed','2 / 42','Modest CPU parallelism and reproducibility']], [1.6,1.2,4.2])
p('The forest handles nonlinear relationships and is the saved validation winner. It is less transparent than one tree, can still overfit or shift under a new population, and cannot infer that changing a feature will retain a customer. Trees do not extrapolate in the same way as a linear model. Its actual test performance is reported separately from the selection decision.')

chapter(7)
h('Decision Tree for churn classification')
p('A Decision Tree repeatedly asks feature questions, such as whether call usage is below a learned split. Each terminal leaf estimates class probabilities from the weighted training examples reaching it. A single tree is easier to follow than a forest but can overfit or change substantially when the sample changes. It is a supervised classifier: Churn labels guide its splits.')
p('The code tries 12 combinations: max_depth [3,5,8], min_samples_split [2,10], and min_samples_leaf [2,5]. All use class_weight balanced and seed 42. Fit each complete pipeline on training only, then choose highest validation F1, with validation ROC-AUC as tie-breaker. The test set never chooses these values.')
cfg=M['Decision Tree']['configuration']
table(['Selected setting','Actual value'],[[key,cfg[key]] for key in ['max_depth','min_samples_split','min_samples_leaf','class_weight','random_state']],[3.5,3.5])
h('K Means for customer segmentation')
p('K-Means starts with k centres, assigns each transformed customer to its nearest centre, and repeatedly updates centres to group means. It minimises within-cluster squared distances, not churn errors. This implementation fits only the ten input features from training. It excludes Churn, identifiers, Status, Customer Value and Age, and uses no class weights.')
p('K-Means uses the same preprocessing definition but fits a separate training-only transformer. Eight numeric or ordinal features are standardised; Complains and Tariff Plan are fully one-hot encoded. A binary category mismatch contributes squared distance 2. Ordinal steps are treated equally. Correlation and encoding can influence distances; groups are not objective customer value.')
S=R['segmentation']
p('Try k=2 to 6 with n_init 10, max_iter 300 and seed 42. Calculate silhouette on the same fixed sample of at most 1,000 training rows. Reject tiny-cluster candidates below 2% of training rows when feasible, then choose the smallest k within 0.01 of the best eligible silhouette. Full-training inertia is shown as an elbow diagnostic, not proof of a unique optimum. No labels or holdout scores choose k.')
p(f'The actual selected k is {S["selected_k"]}; training-sample silhouette is {S["training_silhouette"]:.4f} and full-training inertia is {S["training_inertia"]:,.4f}. Segment IDs are arbitrary, not ordered risk levels. Training-derived names describe usage characteristics, never churn labels.')

chapter(8)
h('Read the training loop')
p('src/train.py loads and prepares the dataset once, builds three classifiers, and wraps each in a Pipeline containing preprocessing followed by the model. pipeline.fit sees training data only. A timer measures preprocessing plus estimator fitting, excluding imports, downloading, explanations and plotting. Each model is scored on validation, and its bundle is saved with joblib.')
p('The winner maximises validation F1 for churn; validation ROC-AUC is a secondary tie-breaker. The threshold is 0.5 for every model. Decision Tree has a bounded 12-configuration validation search. Logistic Regression and Random Forest retain fixed configurations. No early stopping or cross-validation search is claimed.')
table(['Model','Validation F1'],[[name,f'{entry["validation"]["f1"]:.4f}'] for name,entry in M.items()],[4.8,2.2])
h('Why Random Forest remains selected')
p('Random Forest has the highest saved validation F1 among the three classifiers. Decision Tree tuning and classifier selection both use validation F1, with ROC-AUC breaking ties. Selection is frozen before reporting test scores. A stronger score on any future test observation must not silently replace the choice. K-Means is never a competitor in this classifier ranking; it has a different objective and evaluation scale.')
h('What is persisted')
bullet('Each .joblib file contains the fitted preprocessing and estimator, ordered features, schema, threshold, seed, versions, training configuration, excluded-field reasons and dataset SHA-256.')
bullet('results.json holds the audit, partitions, selected name, threshold, model metrics, curves, timing and validation permutation importance.')
bullet('Per-model .joblib.test.csv files store source index, observed target and predicted probability. sample_customers.csv uses real training rows. clean_data.csv contains retained profiles for insights.')
p('Only load trusted local joblib files: loading a malicious pickle-based file can execute code. The UI caches model resources so ordinary interactions reuse the saved model. Package mismatch and missing artifacts produce actionable guidance rather than interactive retraining.')

p('The K-Means bundle separately stores fitted preprocessing, centroids, configuration, feature schema, versions, segment names and training-fitted PCA. segment_assignments.csv holds cluster IDs and 2D coordinates by split. Churn rates are attached only after fitting for descriptive comparison; segment IDs are not fed back into the three churn classifiers.')

chapter(9)
tn,fp=M[R['selected_model']]['test']['confusion_matrix'][0]
fn,tp=M[R['selected_model']]['test']['confusion_matrix'][1]
p('The positive class is churn (1). A true positive is a churner detected; a false negative is a churner missed. A false positive is an alert for a customer who does not churn. A true negative is a correctly recognised non-churner.')
table(['Random Forest test','Predicted non churn','Predicted churn'],[['Actual non churn',tn,fp],['Actual churn',fn,tp]],[2.6,2.2,2.2])
table(['Metric','Formula or meaning','Worked result'],[
 ['Accuracy','(TP + TN) / all rows',f'({tp} + {tn}) / 420 = 0.9476'],['Precision','TP / (TP + FP)',f'{tp} / ({tp} + {fp}) = 0.8281'],
 ['Recall','TP / (TP + FN)',f'{tp} / ({tp} + {fn}) = 0.8281'],['Churn F1','2PR / (P + R)','0.8281'],
 ['ROC AUC','Area under TPR against FPR across score thresholds','0.9851'],['Average precision','Sum of recall increases weighted by precision','0.9200']], [1.3,3.2,2.5])
p('F1 is a harmonic mean, so a weak precision or recall pulls it down. For a small example, suppose a model raises ten alerts and eight are correct: precision is 8/10. If there were sixteen real churners, recall is 8/16. The illustrative F1 is then about 0.615. Those small example values explain the formula; they are not project scores.')
p('ROC-AUC evaluates ranking over thresholds and is not the accuracy at 0.5. TPR equals recall, while FPR = FP / (FP + TN). A random ranking has ROC-AUC near 0.5, subject to finite-sample variation. A good AUC does not prove probability calibration.')
p('PR-AUC in this code means sklearn average_precision_score: AP = sum over thresholds of (R_n - R_previous) times P_n. It is not trapezoidal PR integration. PR curves are useful when churn is uncommon because they directly show how many positive alerts are correct.')
p('Always predicting non-churn would score 356/420 = 84.76% accuracy on this test set, but detect no churners. That is why the project selects using churn F1 and reports several metrics instead of celebrating accuracy alone.')

chapter(10)
p('These values come directly from artifacts/results.json for the revised actual training run. The original split assignments are preserved; model outputs and metrics were regenerated. All three classifiers use the same ten features and threshold 0.5. K-Means uses those inputs without a churn threshold.')
table(['Model','Val F1','Test F1','Test AUC','Fit seconds'],[[n,f'{e["validation"]["f1"]:.4f}',f'{e["test"]["f1"]:.4f}',f'{e["test"]["roc_auc"]:.4f}',f'{e["training_seconds"]:.2f}'] for n,e in M.items()],[2.2,1.1,1.1,1.3,1.3])
table(['Model','Accuracy','Precision','Recall','Avg precision'],[[n,f'{e["test"]["accuracy"]:.4f}',f'{e["test"]["precision"]:.4f}',f'{e["test"]["recall"]:.4f}',f'{e["test"]["pr_auc"]:.4f}'] for n,e in M.items()],[2.2,1.2,1.2,1.1,1.3])
h('What the numbers support')
p('The tree ensembles perform better than the linear baseline on churn F1 in this split. Logistic Regression detects a substantial fraction of churners but produces more false positives, reducing precision. The forest has equal test precision and recall here because its numbers of false positives and false negatives are equal; this is not a property guaranteed by the algorithm.')
p('Decision Tree has high test recall but lower precision than the forest, so it raises more false alerts in this split. Those are descriptive observations, not a universal ordering. Small test samples can change rankings. K-Means has no row in these classifier tables because its silhouette and inertia are not accuracy, F1 or ROC-AUC.')
h('The experiment population')
p(f'Raw class counts are {A["raw_class_counts"]["0"]} non-churn and {A["raw_class_counts"]["1"]} churn. After cleaning they are {A["clean_class_counts"]["0"]} and {A["clean_class_counts"]["1"]}. Test has 420 profiles, including 64 churners. Compare scores only when the data and protocol match; published results on banking or IBM Telco cannot be directly ranked against these results.')
p('Fit times depend on hardware and exclude startup and chart generation. The saved run timestamp is '+R['created_utc']+'. The SHA-256 and row indices in the audit tie the outputs to their source file and split.')

chapter(11)
picture('01_overview.jpg','Figure 1. Final Overview page with real dataset KPIs and customer charts.')
p('Start a demonstration here. Original customers and observed churn use all 3,150 source records; Profiles analysed uses the 2,797 retained profiles. The validation card reports Random Forest’s saved F1, rather than mixing a test value into a selection label.')
p('The left chart counts original churn and non-churn records. The right histogram describes subscription months among retained profiles. Their populations are deliberately labelled because cleaning changes the dataset.')
p('The selected-model panel explains the threshold and frozen selection. How to use provides short routes to single prediction, CSV scoring and evidence. Source and licence remain in the sidebar expander; cleaning details are available below the main content without overwhelming the opening screen.')

chapter(12)
picture('02_single_form.jpg','Figure 2. Final single-customer form with account, usage and service sections.')
p('Choose a saved model; the default is Random Forest. Account and service groups subscription months, tariff, charge band and age band. Usage and engagement groups call seconds, calls, SMS and distinct numbers. Service experience groups failures and complaints. Labels explain units, while help text and category captions retain the documented meaning.')
p('The defaults are taken from the first saved sample training row, not fabricated as a typical customer. Review all ten values and select Estimate churn risk. Numeric fields cannot go below zero, category selectors expose only allowed values, and the existing validator checks the submitted row again.')
p('On a narrow screen, paired columns stack into one reading order. The sidebar collapses behind a navigation control. The form does not change the feature schema, run fit, or save a new model.')

chapter(13)
picture('03_single_result.jpg','Figure 3. Actual default-profile prediction from the saved Decision Tree pipeline.')
p('The result presents estimated churn probability, predicted class, model used and decision threshold. The displayed Decision Tree result for the default sample is 0.0%, below the 50% threshold, so it becomes Non-churn (0). A leaf estimate of zero is not a guarantee that a person cannot churn. The displayed percentage is rounded; the stored probability is used for classification. No arbitrary low, medium or high categories are introduced.')
p('For a what-if demonstration, change an input and submit again. The UI shows the probability change in percentage points when the previous submission used the same model. It avoids comparing deltas across different models. The summary represents the last submitted inputs; edits must be submitted to recalculate it.')
p('If a score changes after reducing failed calls, that means the fitted function responds to that input change. It does not show that fixing failures will cause retention. Model probabilities are uncalibrated estimates from this dataset and class-weighted training, not guarantees for individuals.')

chapter(14)
picture('04_batch_results.jpg','Figure 4. Verified five-row batch result with IDs, probabilities and predicted classes visible first.')
p('Download the sample CSV, keep all ten exact feature headers, and add rows. Churn labels are optional. Extra columns such as customer_id are preserved but never fed into the estimator. Upload is limited to UTF-8 CSV, 5 MB and 50,000 nonempty customer rows. Missing inputs are rejected.')
p('The upload handler calls read_batch and predict rather than implementing a second ML path. Invalid categories, nonnumeric values, infinities, fractional or negative counts, duplicate headers and missing columns produce useful errors. A browser check with Complains removed reported that exact missing column.')
p('After validation, counts show customers scored and predicted churn/non-churn at the saved threshold. The preview puts IDs and predictions first; the export retains original input columns plus the new outputs. Only the preview is limited to 100 rows. Browser verification downloaded five rows and confirmed that all supplied IDs survived. Each row also has customer_segment and segment_profile from the saved K-Means pipeline. These identify an arbitrary usage group rather than a churn class or ordered risk tier.')

chapter(15)
picture('06_model_comparison.jpg','Figure 5. Model comparison shows a selected-model notice, partition choice and uniformly formatted metrics.')
p('The partition control explicitly separates Test final evaluation from Validation model selection. Random Forest is identified using results.json, not hard-coded as a winner inferred from the visible test table. All six metrics display four decimal places, with fit time in seconds.')
p('Metric overview uses the same three classifier colours as ROC and precision-recall curves. The second tab explores ranking tradeoffs without modifying the deployed threshold. Confusion matrices in the third tab show actual outcomes on rows and predictions on columns, with 0 and 1 labelled. The metric glossary is expandable. K-Means evaluation remains in the separate Customer segmentation view.')
p('On a phone-width browser, wide tables and tab strips scroll horizontally. This preserves numeric precision and access to the later columns rather than compressing them into unreadable text. Use validation F1 to explain selection, then discuss test scores honestly as final observations.')

chapter(16)
picture('09_insights_patterns.jpg','Figure 6. Customer patterns view with an attribute selector, observed distributions and outcome relationships.')
p('Customer Insights asks two questions: how is an attribute distributed, and how does it differ by observed outcome? Select an attribute using its human-readable label. Histograms count retained profiles; the box plot shows median and spread. Colours consistently distinguish observed Non-churn and Churn.')
p('A separate Customer segmentation tab shows K-Means results; it is not part of the classifier importance ranking. The patterns page reads artifacts/clean_data.csv through a cached loader. It does not recompute splits or training scores. The displayed population is the cleaned collection of retained profiles, not only the test set and not necessarily a representative sample of a current operator.')
p('An observed relationship can suggest a question for further analysis. For example, a difference in complaint patterns might motivate investigating service quality. It cannot establish that complaints cause churn, because other factors and the data-collection process may explain the association.')

chapter(17)
picture('10_feature_importance.jpg','Figure 7. Saved validation permutation importance for Random Forest, with five-repeat standard deviations.')
p('Select Model feature importance, then a saved model. During training, permutation_importance shuffles one original validation column at a time and measures the change in F1 through the complete pipeline. A larger positive decrease suggests the model relies more on that input for validation performance.')
p('Error bars describe variation over five permutations, not confidence intervals across new training populations. Correlated features can substitute for one another; a low value need not mean a feature has no business relevance. Negative decreases can arise from noise. This is a global model-sensitivity explanation, not a reason for a particular individual’s prediction.')
p('The expandable suggestions are explicitly illustrative business rules: review complaints, investigate repeated failures, or offer a voluntary service check. They are not learned actions, measured ROI or causally validated retention policies. A controlled intervention study would be needed to support those claims.')

chapter(18)
picture('16_segmentation.jpg','Figure 8. Actual K-Means evaluation with training-only selection, cluster sizes and split control.')
S=R['segmentation']
p(f'Selected k is {S["selected_k"]}. Silhouette {S["training_silhouette"]:.4f} uses {S["silhouette_sample_size"]} training rows; inertia {S["training_inertia"]:,.4f} uses the entire training feature matrix. Selected preprocessing plus fit time is {S["training_seconds"]:.3f} seconds. Neither score is classifier F1 or a churn-prediction accuracy.')
table(['Split']+[f'Segment {i}' for i in range(S['selected_k'])],[[split.title()]+[item['size'] for item in meta['profiles']] for split,meta in S['partitions'].items()],[2]+[5/S['selected_k']]*S['selected_k'])
p('Choose Train, Validation or Test to see the corresponding sizes, original-unit feature means and observed churn rates. Those rates are descriptive outcomes attached after fitting; they never chose k. Names describe two large standardised feature deviations from overall training averages. IDs can be permuted without changing the grouping.')
p('The centroid expander shows transformed-space coordinates and explains encoding assumptions. The selection charts show training silhouette and inertia for k=2 to 6. The PCA plot is labelled as a 2D projection for visualisation only, with its training explained-variance ratio. The cluster model was fitted in the full transformed space, not on the two displayed components. Validation/test/uploaded rows use the same saved nearest-centroid assignment and training-fitted PCA.')

chapter(19)
table(['Test area','What it actually checks'],[
 ['Eight range cases','Rejects invalid complaints, charge, age, tariff, negative/fractional counts, infinity and text'],
 ['Schema and missing values','Missing required column fails; missing prediction input fails; training allowance is explicit'],
 ['CSV input handling','Rejects empty content, duplicate headers, oversize and invalid UTF-8; header-only file fails'],
 ['Partition independence','Retained profiles do not overlap; both labels exist; seven charge anomalies are excluded'],
 ['Saved model inference','Every model scores one/many rows consistently; IDs survive; reordered input columns give the same result'],
 ['Preprocessing fit scope','Saved numeric scaler agrees with a fresh fit on the training partition'],
 ['Selection and missing files','Winner agrees with saved validation ranking; missing results give a training command'],
 ['Streamlit smoke test','Every page renders; single prediction and a second what-if submission complete']], [2.1,4.9])
p('The updated regression suite checks three classifiers and K-Means. The final check record in reports/verification.md gives actual pass counts and warnings. Passing these tests is evidence for the checked behaviours, not a proof of all possible inputs, calibration, fairness or cross-platform correctness. For example, the scaler test does not inspect every internal fitted attribute of every transformer.')
p('New segmentation tests verify training-only numeric transforms and centroids, unchanged assignments when supplied churn labels are changed, exact input-feature exclusion, the k-selection rule, one/many assignment agreement, saved split assignments and missing-cluster-artifact errors. Decision Tree tuning is checked against its saved validation candidates.')
h('Browser checks after the revision')
p('All five pages were inspected at desktop 1440 by 1000 and narrow 390 by 844 viewports. Single prediction produced a real result. The sample download, unlabelled five-row upload, CSV export, identifier preservation and missing-column error were exercised. Metric, curve, confusion and importance views were inspected. These are browser emulations of layout, not physical-phone tests.')
p('The original train/validation/test source indices are unchanged. All active models and metrics were regenerated for the revised algorithm inventory; the ten-page PDF is regenerated to match it. Screenshots under reports/ui_screenshots are captures of the final local app. Full check details are in reports/UI_VERIFICATION.md; no cloud deployment is claimed. The new classifier and clustering experiments were actually executed.')

chapter(20)
h('Use the existing environment for the demonstration')
p('Open PowerShell in the project folder. The included artifacts allow inference without training. The tested version is Python 3.13.13 on Windows. Use the exact lock file when loading the included models.')
code(["Set-Location 'C:\\Users\\Ayush Lambat\\OneDrive\\Dokumen\\Coding\\ML'",
      r'.\.venv\Scripts\python.exe -m streamlit run app.py'])
p('If starting on a fresh machine, create a virtual environment and install requirements-lock.txt. No activation-policy change is needed when calling its Python directly.')
code(['python -m venv .venv',r'.\.venv\Scripts\python.exe -m pip install -r requirements-lock.txt'])
h('Reproduce the experiment when needed')
code([r'.\.venv\Scripts\python.exe -m src.data --download',r'.\.venv\Scripts\python.exe -m src.train',
      r'.\.venv\Scripts\python.exe -m pytest -q',r'.\.venv\Scripts\python.exe -m src.report'])
p('Training replaces artifacts and generates charts; only do it when intentionally reproducing or changing the experiment. The ten-page PDF generator uses those saved results. The UI never invokes these commands automatically. If download fails, copy the official CSV to data/raw/iranian_churn.csv. If venv cannot bootstrap pip, the README gives a system-pip --python fallback.')
h('Deployment instructions not completed deployment')
p('Review the project and create your GitHub repository, then push code, attribution and trusted artifacts when authorised. In Streamlit Community Cloud choose that repository, branch and app.py entrypoint. Select Python 3.13 in Advanced settings to match the artifacts. Use the exact locked dependencies for reproducibility, or retrain in the chosen compatible environment.')
p('After deploying, inspect build logs and repeat single prediction, upload, download and chart checks. Add the genuine URL and screenshots to submission materials. No remote push, public app or deployment URL is claimed here. Locally, add --server.port 8502 if the standard port is occupied; this redesign was reviewed on 8502.')
p('Missing artifacts: restore them or run training. Version mismatch: install the lock file or intentionally retrain, rather than ignoring the check. For macOS/Linux the README substitutes .venv/bin/python; those installations have not been tested in this environment.')

chapter(21)
h('What this project cannot establish')
bullet('Generalisability: one historical Iranian telecom file cannot establish performance for every operator, country or later time period.')
bullet('Independence and timing: missing IDs and original timestamps limit the leakage audit. Collapsing profiles can merge different people and alters class balance.')
bullet('Schema certainty: seven charge=10 records conflict with the notes. Exclusion is documented, but the true intended meaning remains unresolved.')
bullet('Statistical certainty: one split is reported, without confidence intervals. A validation winner may not have the best observed test score.')
bullet('Probability meaning: class weighting and lack of calibration mean an 80% score need not correspond to an 80% observed event frequency.')
bullet('Fairness and causality: age groups are used without a subgroup fairness study. Associations, importance and what-if changes do not measure intervention effects.')
bullet('Operational delivery: no physical-device, macOS/Linux or public-cloud validation was completed. Local browser review is a narrower claim.')
h('Useful next improvements')
p('First resolve feature timing and charge-band definitions with the data provider. Then obtain timestamped, person-identified data under an appropriate privacy process and evaluate on a later period or external operator. Track performance decay and data drift rather than relying on a one-time score.')
p('For probability use, assess calibration on a separate development partition and examine subgroup errors. Repeated or cross-validated development experiments can describe variation, but preserve a fresh final holdout. Tune an operating threshold only on development data, guided by actual costs of false alerts and missed churners.')
p(R['rubric_note'])
p('For retention claims, test a voluntary business intervention with a suitable controlled design. Measure both benefit and unwanted effects. Keep the model’s ranking task separate from the causal question of which offer helps which customer. A more complex model alone does not resolve those evidence gaps.')

chapter(22)
qa=[
 ('What is churn in this project?','The binary target is 1 for churn and 0 for non-churn. Inputs describe the first nine months; UCI states the outcome is measured at twelve months.'),
 ('Why are there ten inputs if UCI lists thirteen predictors?','Status, Customer Value and the separately unexplained Age column are excluded. The ten documented retained inputs are explicitly listed in config.py.'),
 ('Why remove seven charge records?','Their value 10 conflicts with the published 0-9 definition. Raw data is preserved and indices are saved. They may be undocumented valid bands; we do not know for certain.'),
 ('Why not demonstrate cleaning by deleting big values?','Large nonnegative counts can be valid. Cleaning should follow evidence and documented constraints, not remove observations just to make a chart look tidy.'),
 ('Why three partitions?','Training fits transformations and estimators, validation selects the model, and test gives final reporting after the choice is frozen.'),
 ('Why is the split stratified?','Churn is a minority class. Stratification keeps both classes represented in similar proportions across the partitions.'),
 ('Why does Random Forest remain selected?','It has the highest validation F1 among Logistic Regression, Random Forest and the tuned Decision Tree. The choice is frozen before test reporting. K-Means has a different unsupervised objective and is not ranked against their F1 scores.'),
 ('What is the main difference between a forest and a single tree?','A single Decision Tree follows one set of branching rules. Random Forest averages many randomised trees, which can reduce variance. Both learn churn labels in this project.'),
 ('Is a K-Means cluster number a churn prediction?','No. It is an arbitrary nearest-centroid group ID. K-Means does not learn Churn, and a larger cluster number is not a higher risk level.'),
 ('How do new inputs receive the same scaling?','The saved Pipeline contains the already fitted imputer, scaler and encoder. predict_proba transforms new values before classification; it does not refit.'),
 ('Why is accuracy insufficient?','Most profiles are non-churn. Predicting non-churn for every test row would give 84.76% accuracy but zero churn recall.'),
 ('What does PR-AUC mean here?','Average precision from sklearn: recall increments weighted by precision. It is not trapezoidal PR integration.'),
 ('What exactly is the threshold?','A probability at least 0.5 becomes churn. The threshold was fixed in advance and is displayed beside the prediction.'),
 ('Do what-if predictions tell us how to retain someone?','No. They show how a fitted model responds to changed inputs. Retention effects need causal evidence or a controlled study.'),
 ('Can silhouette be compared with F1?','No. Silhouette measures feature-distance cohesion and separation; F1 measures labelled churn precision and recall. The project has three classifiers and one clustering model, which may need instructor approval for a four-classifier rubric.'),
]
for question,answer in qa:
    para=p(); para.paragraph_format.keep_with_next=True
    run=para.add_run(question); run.bold=True
    p(answer)

chapter(23)
h('Primary data and deployment sources')
p('UCI Machine Learning Repository. Iranian Churn [Dataset] (2020). DOI 10.24432/C5JW3Z. CC BY 4.0. https://archive.ics.uci.edu/dataset/563/iranian+churn+dataset')
p('Streamlit. Deploy your app on Community Cloud. https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/deploy. The README records the checked deployment guidance; deployment remains a future instruction.')
h('Research behind the original submission')
papers=json.loads((REPORTS/'literature.json').read_text(encoding='utf-8'))
for i,paper in enumerate(papers,1):
    p(f'{i}. {paper["authors"]} ({paper["year"]}). {paper["title"]}. DOI {paper["doi"]}. {paper["url"]}')
p('The revised literature review uses five verified recent primary publisher or author-university records relevant to classification, preprocessing and segmentation. Reviews are focused comparisons, not a full-text systematic review. Scores and cluster counts from different datasets are not directly comparable with this experiment.')
h('Project evidence to consult during a viva')
p('artifacts/results.json provides the source audit, split indices, scores, selection and importance. src/train.py, preprocessing.py, data.py, evaluate.py and predict.py establish the actual ML path. tests/ contains executable checks. reports/UI_VERIFICATION.md records the completed redesign review. README.md provides setup, deployment and troubleshooting. The regenerated Customer_Churn_Report.pdf is the separate ten-page submission report and includes the classification-versus-clustering rubric clarification.')

doc.save(OUT)
print(f'Created {OUT}; {len(TITLES)} chapters; {len(doc.inline_shapes)} actual UI screenshots; verify page count after rendering')

if __name__=='__main__':
    pass
